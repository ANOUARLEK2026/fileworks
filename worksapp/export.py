"""تصدير الوثائق إلى Word وPDF (python-docx، يعمل دون إنترنت).

* الوثيقة الصادرة تُصدَّر من اللقطة المحفوظة فقط (لا من البيانات الحية).
* المسودة تحمل شريط «مسودة — غير معتمدة»، والنسخة التصحيحية تحمل إشارة التصحيح وسببه.
* التصدير منفصل عن الاعتماد: لا يغيّر حالة الوثيقة أبدًا.
* لا تُكتب فوق ملفات النماذج الأصلية: الملف الناتج جديد دائمًا.

الكميات تُعرض حسب الوحدة ونوع الوثيقة (qtyfmt)، والمبلغ بالحروف من numwords (أو arabic_money بإعداد words_style)،
وبيانات الشركة المتعاقدة من لقطة الشركة (companies) لا من المشروع، مع الرجوع لحقول المشروع القديمة للوثائق الصادرة سابقًا.
PDF يعتمد على LibreOffice المثبت محليًا؛ لا تُرفع الملفات إلى أي خدمة.
"""
import json
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from . import services as SV, storage as S
from .money import D, fmt as _fmt, bidi_safe
from .numwords import amount_in_words
from .qtyfmt import format_qty
from . import arabic_money


def fmt(v, places=2):
    """فاصل الآلاف بمسافة غير قابلة للكسر كي لا يعيد الاتجاه RTL ترتيب مجموعات الرقم."""
    return _fmt(v, places).replace(" ", "\u00a0")


def ltr(s):
    """يثبّت ترتيب التواريخ والأرقام ذات الفواصل داخل نص RTL (علامة LRM حول - / : .) كي لا تنعكس مقاطعها."""
    if s in (None, ""):
        return s
    return re.sub(r"([-/:.])", "\u200e\\1\u200e", str(s))


def factor(x):
    """عامل قياس كما يُكتب في النموذج: 0.25 لا .25 ، وبمنزلتين على الأقل (0.60)."""
    if x in (None, ""):
        return ""
    try:
        txt = format(D(x), "f")
    except Exception:
        return str(x)
    if "." not in txt:
        return txt
    head, tail = txt.split(".")
    return f"{head}.{tail.ljust(2, '0')}"

FONT = "Arial"
# محاذاة دلالية: R = «بداية السطر» (أقصى اليمين في RTL)، C = وسط. في فقرة bidi يفسّر Word وLibreOffice القيمة
# jc=left على أنها بداية السطر (اليمين) وjc=right على أنها نهايته (اليسار)؛ لذلك نكتب العنصر مباشرة عبر _jc
# بدل WD_ALIGN_PARAGRAPH.RIGHT الذي كان يُلصق النص باليسار فعليًا.
R, C = "start", "center"
_JC_BIDI = {"start": "left", "end": "right", "center": "center", "both": "both"}


def _jc(p, align):
    ppr = p._p.get_or_add_pPr()
    for old in ppr.findall(qn("w:jc")):
        ppr.remove(old)
    el = OxmlElement("w:jc")
    el.set(qn("w:val"), _JC_BIDI.get(align, "left"))
    ppr.append(el)


def _bidi(p):
    ppr = p._p.get_or_add_pPr()
    if ppr.find(qn("w:bidi")) is None:
        ppr.append(OxmlElement("w:bidi"))


# ---------- أدوات RTL ----------
def _set(parent, tag, **attrs):
    el = parent.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        parent.append(el)
    for k, v in attrs.items():
        el.set(qn(k), v)
    return el


_RPR = ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline", "shadow", "emboss",
        "imprint", "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing", "w", "kern", "position", "sz", "szCs",
        "highlight", "u", "effect", "bdr", "shd", "fitText", "vertAlign", "rtl", "cs", "em", "lang"]
_TBLPR = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW", "jc",
          "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook"]


def _reorder(el, order):
    def key(c):
        name = c.tag.split("}")[1]
        return order.index(name) if name in order else len(order)
    kids = sorted(list(el), key=key)
    for k in list(el):
        el.remove(k)
    for k in kids:
        el.append(k)


def run(p, text, bold=False, size=11, underline=False, color=None):
    r = p.add_run(str(text))
    r.font.name = FONT
    r.font.size = Pt(size)
    r.bold = bold
    r.underline = underline
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    rpr = r._r.get_or_add_rPr()
    _set(rpr, "w:rFonts", **{"w:ascii": FONT, "w:hAnsi": FONT, "w:cs": FONT})
    _set(rpr, "w:rtl")
    _set(rpr, "w:szCs", **{"w:val": str(int(size * 2))})       # النص العربي (معقّد) يستعمل szCs/bCs
    if bold:
        _set(rpr, "w:bCs")
    _reorder(rpr, _RPR)
    return r


def para(container, text="", bold=False, size=11, align=R, underline=False, space_after=4, color=None):
    p = container.add_paragraph()
    _bidi(p)
    _jc(p, align)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    if text != "":
        run(p, text, bold, size, underline, color)
    return p


def _cell_text(cell, text, bold=False, size=9, align=C, shade=None):
    cell.text = ""
    p = cell.paragraphs[0]
    _bidi(p)
    _jc(p, align)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.space_before = Pt(1)
    t = "" if text is None else text
    if isinstance(t, str) and t and re.fullmatch(r"[\s\u00a0\d.,\-—/:%+*]+", t) and any(ch.isdigit() for ch in t):
        t = bidi_safe(t)
    run(p, t, bold, size)
    if shade:
        tcpr = cell._tc.get_or_add_tcPr()
        for old in tcpr.findall(qn("w:shd")):          # لا نكرر shd عند إعادة كتابة الخلية
            tcpr.remove(old)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), shade)
        tcpr.append(shd)


def new_table(doc, ncols, widths_cm):
    t = doc.add_table(rows=0, cols=ncols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tblpr = t._tbl.tblPr
    _set(tblpr, "w:bidiVisual")                       # العمود الأول يظهر في أقصى اليمين
    _set(tblpr, "w:tblLayout", **{"w:type": "fixed"})
    for i, w in enumerate(widths_cm):
        t.columns[i].width = Cm(w)
    _reorder(tblpr, _TBLPR)
    t._widths = widths_cm
    return t


def add_row(t, values, bold=False, size=9, shade=None, header=False, aligns=None):
    row = t.add_row()
    trpr = row._tr.get_or_add_trPr()
    _set(trpr, "w:cantSplit")
    if header:
        _set(trpr, "w:tblHeader")                     # تتكرر رؤوس الجداول في الصفحات التالية
    for i, v in enumerate(values):
        row.cells[i].width = Cm(t._widths[i])
        al = aligns[i] if aligns else C
        _cell_text(row.cells[i], v, bold or header, size, al, shade or ("E6ECF0" if header else None))
    return row


def keep_together(t):
    """يمنع انقسام جدول صغير (ملخص/حوصلة) بين صفحتين: كل سطر يلتصق بما بعده."""
    for row in t.rows[:-1]:
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.keep_with_next = True


def merge(row, a, b, text, bold=True, size=9, align=R, shade=None):
    m = row.cells[a].merge(row.cells[b])
    _cell_text(m, text, bold, size, align, shade)
    return m


def page(doc, landscape):
    s = doc.sections[0]
    if landscape:
        s.orientation = WD_ORIENT.LANDSCAPE
        s.page_width, s.page_height = Cm(29.7), Cm(21.0)
    else:
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(1.3)
    s.top_margin = s.bottom_margin = Cm(1.3)
    return s


# ---------- نموذج الوثيقة ----------
def overlay_company(state):
    """يدمج بيانات الشركة المتعاقدة (من لقطة الشركة) في نسخة من المشروع: contractor + حقول المؤسسة في extra_json،
    فتقرأها دوال البناء من مكان واحد. اللقطات القديمة بلا شركة تبقى كما هي (حقول المشروع القديمة)."""
    co = state.get("company")
    if not co:
        return state
    proj = dict(state["project"])
    extra = json.loads(proj.get("extra_json") or "{}")
    for k in S.COMPANY_KEYS:
        if co.get(k):                      # قيمة الشركة تغلب؛ وحقول المشروع القديمة تبقى احتياطًا لقواعد ما قبل الفصل
            extra[k] = co[k]
    proj["contractor"] = co.get("name") or proj.get("contractor")
    proj["extra_json"] = json.dumps(extra, ensure_ascii=False)
    return {**state, "project": proj}


def load_model(conn, doc_id):
    d = S.get_document(conn, doc_id)
    if d["status"] == "issued":
        snap = json.loads(d["snapshot_json"])
        return d, overlay_company(snap["state"]), snap["results"]
    c = SV.compute(conn, doc_id)
    return d, overlay_company(S.capture_project_state(conn, d["project_id"])), c["results"]


def _num(d):
    """رقم الوثيقة، مع «والنهائي» إن اختار المستخدم ذلك (كما في عنوان النموذج: «رقم: 01 والنهائي»)."""
    base = "(مسودة)" if d["number"] is None else f"{d['number']:02d}"
    final = json.loads(d.get("payload_json") or "{}").get("config", {}).get("final")
    return base + (" والنهائي" if final else "")


def _summary_style(state):
    return json.loads(state["project"].get("settings_json") or "{}").get("summary_style") or "simple"


def _banner(doc, d):
    if d["status"] == "draft":
        para(doc, "مسودة — غير معتمدة", True, 11, C, color="B3402F")
    if d["revision"]:
        para(doc, f"نسخة تصحيحية رقم {d['revision']} — السبب: {d['correction_reason'] or ''}", True, 10, C, color="B3402F")


def _letterhead(doc, proj):
    extra = json.loads(proj.get("extra_json") or "{}")
    para(doc, proj.get("contractor") or "", True, 15, C, space_after=2)
    if extra.get("manager"):
        para(doc, f"المسير: {extra['manager']}", True, 11, C, space_after=2)
    if extra.get("address"):
        para(doc, extra["address"], False, 10, C, space_after=2)


def _project_lines(doc, proj):
    para(doc, f"المشروع: {proj.get('name') or ''}", True, 11)
    if proj.get("contract_no") or proj.get("contract_date"):
        para(doc, f"عقد رقم: {ltr(proj.get('contract_no')) or ''}      بتاريخ: {ltr(proj.get('contract_date')) or ''}", True, 10)


def _signatures(doc, labels, names=None, place_date=None):
    if place_date:
        para(doc, place_date, True, 10, space_after=4)
    t = new_table(doc, len(labels), [6.0] * len(labels))
    row = t.add_row()
    for i, lab in enumerate(labels):
        _cell_text(row.cells[i], "", size=11)
        p = row.cells[i].paragraphs[0]
        run(p, lab, True, 10, underline=True)
        if names and names[i]:
            p2 = row.cells[i].add_paragraph(); _bidi(p2); _jc(p2, C)
            run(p2, names[i], True, 10)
    for c in row.cells:
        tcpr = c._tc.get_or_add_tcPr()
        b = OxmlElement("w:tcBorders")
        for side in ("top", "left", "bottom", "right"):
            e = OxmlElement(f"w:{side}"); e.set(qn("w:val"), "nil"); b.append(e)
        tcpr.append(b)
    para(doc, "", space_after=40)


def _items(state):
    return {i["code"]: i for i in state["items"]}


def _qf(state, doc_type):
    """دالة عرض الكمية للوثيقة: q(قيمة، وحدة). القواعد الافتراضية في qtyfmt وتُجاوَز من إعداد المشروع qty_format."""
    over = json.loads(state["project"].get("settings_json") or "{}").get("qty_format")
    return lambda v, unit=None: format_qty(v, unit, doc_type, over)


def _cfg(d):
    return json.loads(d.get("payload_json") or "{}").get("config", {})


def _extra(proj):
    return json.loads(proj.get("extra_json") or "{}")


def _place_date(d, proj):
    """«الوادي في: 2026-03-31» — لا يُطبع شيء إن لم يُدخل المكان والتاريخ."""
    c, e = _cfg(d), _extra(proj)
    place, date = c.get("place") or e.get("place"), ltr(c.get("date"))
    if date and place:
        return f"{place} في: {date}"
    return f"التاريخ: {date}" if date else (f"{place} في:" if place else None)


def _words_line(doc, prefix, amount, state=None):
    """سطر «… بالحروف». المبلغ غير المحسوب لا يُكتب (لا نفترض صفرًا).
    الأسلوب الافتراضي يطابق النماذج الممسوحة؛ وعند words_style=formal يكتب «فقط … لا غير»."""
    style = None
    if state is not None:
        style = json.loads(state["project"].get("settings_json") or "{}").get("words_style")
    w = arabic_money.amount_in_words(amount) if style == "formal" else amount_in_words(amount)
    if w:
        para(doc, f"{prefix} {w}", True, 10, space_after=6)


# ---------- بيان الأشغال ----------
def build_bayan(doc, d, state, res):
    proj = state["project"]
    _q = _qf(state, "bayan")
    page(doc, False)
    _banner(doc, d)
    _letterhead(doc, proj)
    para(doc, "", space_after=30)
    box = new_table(doc, 1, [12.0])
    row = box.add_row()
    _cell_text(row.cells[0], "", size=16)
    p = row.cells[0].paragraphs[0]
    run(p, "بيان الأشغال المنجزة", True, 20)
    p2 = row.cells[0].add_paragraph(); _bidi(p2); _jc(p2, C)
    run(p2, f"رقم: {_num(d)}", True, 16)
    para(doc, "", space_after=24)
    _project_lines(doc, proj)
    doc.add_page_break()

    para(doc, f"المشروع: {proj.get('name') or ''}", True, 10)
    para(doc, f"بيان الأشغال المنجزة رقم: {_num(d)}", True, 12, C, space_after=8)
    for r in res["rows"]:
        para(doc, f"{r['code']}   {r.get('description') or ''}", False, 10, space_after=2)
        if r["mode"] == "measured" and r["lines"]:
            t = new_table(doc, 3, [5.0, 8.0, 4.0])
            for ln in r["lines"]:
                add_row(t, [ln["location"] or ln["label"] or "", "  ×  ".join(factor(x) for x in ln["factors"] if x != ""),
                            "= " + _q(ln["result"], r.get("unit"))], size=9)
        tot = new_table(doc, 2, [6.0, 4.0])
        txt = f"{_q(r['qty'], r.get('unit'))} {r.get('unit') or ''}"
        add_row(tot, ["المجموع الكلي", txt], bold=True, size=10)
        if r.get("approved"):
            para(doc, f"(كمية معتمدة بعد حسم التعارض من: {r['approved_source']}؛ السبب: {r['approved_reason']})",
                 False, 8, color="285F88")
        elif r["overridden"]:
            para(doc, f"(عُدّلت الكمية يدويًا؛ المحسوبة: {_q(r['computed'], r.get('unit'))})", False, 8, color="B3402F")
        para(doc, "", space_after=6)
    _signatures(doc, ["المتعامل المتعاقد", "مكتب الدراسات"], [_extra(proj).get("manager"), None], _place_date(d, proj))


# ---------- حوصلة الأشغال ----------
def build_hawsala(doc, d, state, res):
    proj = state["project"]
    _q = _qf(state, "hawsala")
    page(doc, False)
    _banner(doc, d)
    _letterhead(doc, proj)
    para(doc, f"حوصلة الأشغال المنجزة رقم: {_num(d)}", True, 14, C, space_after=6)
    para(doc, f"المشروع: {proj.get('name') or ''}", True, 11, space_after=8)
    t = new_table(doc, 7, [1.2, 7.2, 1.3, 2.0, 2.0, 2.0, 2.2])
    add_row(t, ["رقم", "تعيين الأشغال", "الوحدة", "الكمية المقررة في العقد", "الكمية المنجزة سابقًا",
                "الكمية المنجزة خلال الشهر", "الكمية الكلية المنجزة"], header=True, size=8)
    for r in res["rows"]:
        add_row(t, [r["code"], r.get("description") or "", r.get("unit") or "", _q(r["contract_qty"], r.get("unit")),
                    _q(r["prev"], r.get("unit")), _q(r["period"], r.get("unit")), _q(r["cum"], r.get("unit"))], aligns=[C, R, C, C, C, C, C])
    para(doc, "", space_after=14)
    _signatures(doc, ["المتعامل المتعاقد", "مكتب الدراسات"], [_extra(proj).get("manager"), None], _place_date(d, proj))


# ---------- وضعية الأشغال ----------
def _situation_form_summary(doc, tt, res, rate):
    """ملخص على نسق نموذج الوضعية: (1) تراكمي، (2) سابق، (3)=(1)−(2)، استردادات، (4)، الضمان (5)، الصافي (6).
    عرض فقط لنفس القيم المحسوبة: الصافي هنا = الصافي في الحساب (TTC الفترة − مجموع الاقتطاعات)."""
    ttc_c, ttc_p = D(tt["ttc_cumulative"]), D(tt["ttc_period"])
    prev = ttc_c - ttc_p if ttc_c is not None and ttc_p is not None else None
    ded = [(x["label"], D(x["amount"]), x["code"] == "guarantee") for x in res["deductions"]]
    other = [x for x in ded if not x[2]]
    guar = [x for x in ded if x[2]]
    sum_other = sum((x[1] for x in other), D(0)) if all(x[1] is not None for x in other) else None
    sum_guar = sum((x[1] for x in guar), D(0)) if all(x[1] is not None for x in guar) else None
    v4 = ttc_p - sum_other if ttc_p is not None and sum_other is not None else None
    v6 = v4 - sum_guar if v4 is not None and sum_guar is not None else None
    s = new_table(doc, 2, [17.0, 6.0])
    add_row(s, ["البيان", "المبلغ (دج)"], header=True)
    add_row(s, ["(1) مبلغ الأشغال المنجزة (تراكمي، بكل الرسوم)", fmt(ttc_c)], aligns=[R, C])
    add_row(s, ["(2) يطرح: مبلغ الأشغال المنجزة سابقًا (بكل الرسوم)", fmt(prev)], aligns=[R, C])
    add_row(s, ["(3) المبلغ الخام للوضعية = (1) − (2)", fmt(ttc_p)], bold=True, aligns=[R, C])
    for label, amt, _ in other:
        add_row(s, [f"استردادات / اقتطاعات: {label}", fmt(amt)], aligns=[R, C])
    add_row(s, ["(4) المبلغ الصافي للوضعية = (3) − مجموع الاستردادات", fmt(v4)], bold=True, aligns=[R, C])
    for label, amt, _ in guar:
        add_row(s, [f"(5) {label}", fmt(amt)], aligns=[R, C])
    add_row(s, ["(6) المبلغ الصافي للدفع = (4) − (5)", fmt(v6)], bold=True, shade="EEF2F4", size=10, aligns=[R, C])
    keep_together(s)


def _legal_block(doc, d, proj):
    """بيانات المؤسسة والعقد كما في صفحة الوضعية الأولى. تُطبع الحقول المُدخلة فقط."""
    e = _extra(proj)
    para(doc, "الجمهورية الجزائرية الديمقراطية الشعبية", True, 11, C, space_after=2)
    if e.get("wilaya"):
        para(doc, f"ولاية {e['wilaya']}", True, 10, space_after=2)
    rc = " ".join(x for x in (e.get("rc_no"), f"بتاريخ: {ltr(e['rc_date'])}" if e.get("rc_date") else "") if x)
    for label, val in (("السجل التجاري رقم", rc), ("الرقم الجبائي", e.get("nif")), ("رقم المادة", e.get("article_no")),
                       ("رقم التعريف الإحصائي", e.get("nis")),
                       ("رقم الحساب البنكي", " ".join(x for x in (e.get("rib"), e.get("bank")) if x))):
        if val:
            para(doc, f"{label}: {val}", False, 10, space_after=1)
    para(doc, "1- جزء خاص بالمؤسسة", True, 10, space_after=2)
    who = " — ".join(x for x in (proj.get("contractor"), f"المسير: {e['manager']}" if e.get("manager") else "", e.get("address")) if x)
    for label, val in (("المؤسسة (المقر الاجتماعي والعنوان)", who), ("العقد رقم", ltr(proj.get("contract_no"))),
                       ("المؤرخ في", ltr(proj.get("contract_date"))), ("موضوع العقد", proj.get("name")),
                       ("مبلغ العقد", f"{e['contract_amount']} دج" if e.get("contract_amount") else "")):
        if val:
            para(doc, f"{label}: {val}", False, 10, space_after=1)


def build_situation(doc, d, state, res):
    proj, items = state["project"], _items(state)
    _q = _qf(state, "situation")
    page(doc, True)
    _banner(doc, d)
    _legal_block(doc, d, proj)
    date = ltr(_cfg(d).get("date"))
    para(doc, f"وضعية الأشغال رقم: {_num(d)}" + (f"   بتاريخ: {date}" if date else ""), True, 14, C, space_after=2)
    para(doc, f"المشروع: {proj.get('name') or ''}", True, 11, space_after=6)
    w = [1.1, 7.6, 1.6, 1.9, 1.9, 1.9, 2.0, 2.3, 3.2, 3.2]
    t = new_table(doc, 10, w)
    add_row(t, ["رقم", "تعيين الأشغال", "الوحدة", "الكمية المقررة في العقد", "الكمية سابقًا", "الكمية خلال الشهر",
                "الكمية الكلية المنجزة", "سعر الوحدة", "مبلغ الأشغال المنجزة خلال الشهر", "المبلغ الكلي"], header=True, size=8)
    for l in res["lines"]:
        it = items.get(l["code"], {})
        u = it.get("unit")
        add_row(t, [l["code"], it.get("description") or "", u or "", _q(D(it.get("contract_qty")), u), _q(D(l["qty_prev"]), u),
                    _q(D(l["qty_period"]), u), _q(D(l["qty_cum"]), u), fmt(D(l["price"])), fmt(D(l["amount_period"])),
                    fmt(D(l["amount_cum"]))], aligns=[C, R, C, C, C, C, C, C, C, C])
    tt = res["totals"]
    rate = D(res.get("tva_rate"))
    for label, a, b in (("المجموع العام بدون رسوم (HT)", tt["ht_period"], tt["ht_cumulative"]),
                        (f"الرسم على القيمة المضافة (TVA) {fmt(rate * 100) if rate is not None else '—'}%", tt["tva_period"], tt["tva_cumulative"]),
                        ("المجموع العام بكل الرسوم (TTC)", tt["ttc_period"], tt["ttc_cumulative"])):
        row = add_row(t, [""] * 10, bold=True)
        merge(row, 0, 7, label)
        _cell_text(row.cells[8], fmt(D(a)), True, 9); _cell_text(row.cells[9], fmt(D(b)), True, 9)
    para(doc, "", space_after=8)
    para(doc, "ملخص الحساب", True, 12, space_after=4).paragraph_format.keep_with_next = True
    names = [_extra(proj).get("manager"), None, _extra(proj).get("owner_rep") or proj.get("owner")]
    if _summary_style(state) == "form":
        _situation_form_summary(doc, tt, res, rate)
        para(doc, "", space_after=8)
        _words_line(doc, "المبلغ الصافي للدفع بما فيه الرسم (بالأحرف):", res.get("net"), state)
        _signatures(doc, ["المتعامل المتعاقد", "مكتب الدراسات", "المصلحة المتعاقدة"], names, _place_date(d, proj))
        return
    s = new_table(doc, 3, [11.0, 5.5, 8.0])
    add_row(s, ["البيان", "المبلغ (دج)", "طريقة الحساب"], header=True, size=9)
    add_row(s, ["المجموع بدون رسوم (HT)", fmt(D(tt["ht_period"])), "الفترة"], size=9, aligns=[R, C, R])
    add_row(s, [f"TVA {fmt(rate * 100) if rate is not None else ''}%", fmt(D(tt["tva_period"])), "على HT قبل الاقتطاعات"], size=9, aligns=[R, C, R])
    add_row(s, ["المجموع بكل الرسوم (TTC)", fmt(D(tt["ttc_period"])), ""], bold=True, size=9, aligns=[R, C, R])
    for x in res["deductions"]:
        add_row(s, [x["label"], fmt(D(x["amount"])), x["explain"].split(":", 1)[-1].strip()], size=8, aligns=[R, C, R])
    add_row(s, ["المبلغ الصافي للدفع", fmt(D(res["net"])), ""], bold=True, size=10, shade="EEF2F4", aligns=[R, C, R])
    keep_together(s)
    para(doc, "", space_after=8)
    _words_line(doc, "المبلغ الصافي للدفع بما فيه الرسم (بالأحرف):", res.get("net"), state)
    _signatures(doc, ["المتعامل المتعاقد", "مكتب الدراسات", "المصلحة المتعاقدة"], names, _place_date(d, proj))


# ---------- الجرد العام والنقصان ----------
def _row_amounts(r, included):
    price = D(r["price"])
    inc_amt = r["increase_amount_by_status"]
    inc_final = sum(((inc_amt[s] or D(0)) for s in included), D(0)) if included is not None else None
    shown_inc = inc_final if inc_final is not None else (inc_amt["approved"] or D(0)) + (inc_amt["pending"] or D(0))
    ca, da = D(r["contract_amount"]), D(r["decrease_amount"])
    final = None if inc_final is None or ca is None else ca - da + inc_final
    return ca, da, shown_inc, final


def _inv_norm(res):
    """اللقطة الصادرة تخزّن الأعداد نصوصًا: نعيدها Decimal قبل أي حساب عرض."""
    dd = lambda x: D(x)
    for r in res["rows"]:
        for k in ("contract_qty", "decrease", "executed", "price", "contract_amount", "decrease_amount"):
            r[k] = dd(r.get(k))
        for k in ("increase_by_status", "increase_amount_by_status"):
            r[k] = {a: dd(b) for a, b in r[k].items()}
    t = res["totals"]
    for k in ("ht", "tva"):
        if t.get(k):
            t[k] = {a: dd(b) for a, b in t[k].items()}
    for k in ("final_contractual_ht", "final_contractual_ttc"):
        if k in t:
            t[k] = dd(t[k])
    return res


def build_inventory(doc, d, state, res):
    res = _inv_norm(res)
    proj = state["project"]
    _q = _qf(state, "inventory")
    page(doc, True)
    _banner(doc, d)
    para(doc, f"{proj.get('contractor') or ''}", True, 11, space_after=2)
    para(doc, "الجرد العام والنهائي", True, 15, C, space_after=2)
    para(doc, f"المشروع: {proj.get('name') or ''}      عقد رقم: {ltr(proj.get('contract_no')) or ''}   بتاريخ: {ltr(proj.get('contract_date')) or ''}",
         True, 10, space_after=6)
    # أعمدة المبالغ واسعة بما يكفي لمنع التفاف آخر خانة عشرية داخل الخلية.
    w = [1.2, 3.2, 1.3, 1.7, 1.5, 1.5, 1.7, 2.7, 3.1, 2.9, 3.0, 3.3]   # المجموع 27.1 سم = عرض الصفحة الأفقية بين الهوامش
    t = new_table(doc, 12, w)
    h1 = add_row(t, [""] * 12, header=True)
    for i, label in ((0, "رقم"), (1, "تعيين الأشغال"), (2, "الوحدة"), (3, "الكمية المقررة في العقد"), (6, "الكمية المنجزة كليًا"),
                     (7, "سعر الوحدة"), (8, "مبلغ العقد"), (11, "المبلغ الكلي")):
        _cell_text(h1.cells[i], label, True, 8, C, "E6ECF0")
    merge(h1, 4, 5, "الكمية", True, 8, C, "E6ECF0"); merge(h1, 9, 10, "المبلغ", True, 8, C, "E6ECF0")
    h2 = add_row(t, [""] * 12, header=True)
    for i, label in ((4, "بالنقصان"), (5, "بالزيادة"), (9, "بالنقصان"), (10, "بالزيادة")):
        _cell_text(h2.cells[i], label, True, 7, C, "E6ECF0")
    for i in (0, 1, 2, 3, 6, 7, 8, 11):
        a, b = t.cell(0, i), t.cell(1, i)
        a.merge(b)

    tot = res["totals"]
    included = tot.get("included_statuses")
    sec_cur, sec_n, sub = None, 0, [D(0)] * 4
    pending_note = False

    def flush_subtotal():
        row = add_row(t, [""] * 12, bold=True)
        merge(row, 0, 7, f"المجموع الفرعي {sec_n:02d}")
        for k, col in enumerate((8, 9, 10, 11)):
            _cell_text(row.cells[col], fmt(sub[k]) if (k < 3 or included is not None) else "—", True, 9)

    for r in res["rows"]:
        sec = r.get("section")
        if sec != sec_cur:
            if sec_cur is not None or sec_n:
                flush_subtotal()
            sec_cur, sub = sec, [D(0)] * 4
            if sec:
                sec_n += 1
                row = add_row(t, [""] * 12, bold=True, shade="F3F5F7")
                merge(row, 0, 11, sec, align=C, shade="F3F5F7")
        ca, da, ia, fin = _row_amounts(r, included)
        inc_q = (r["increase_by_status"]["approved"] + r["increase_by_status"]["pending"])
        star = "*" if r["increase_by_status"]["pending"] > 0 else ""
        pending_note |= bool(star)
        u = r.get("unit")
        add_row(t, [r["code"], r.get("description") or "", u or "", _q(r["contract_qty"], u), _q(r["decrease"], u),
                    _q(inc_q, u) + star, _q(r.get("executed"), u), fmt(D(r["price"])), fmt(ca), fmt(da), fmt(ia),
                    fmt(fin) if fin is not None else "—"], aligns=[C, R] + [C] * 10)
        for k, v in enumerate((ca, da, ia, fin)):
            if v is not None:
                sub[k] += v
    if res["rows"] and (sec_cur is not None or sec_n):
        flush_subtotal()

    ht, tva = tot.get("ht"), tot.get("tva")
    if ht and tot.get("final_contractual_ht") is not None:
        inc_shown = sum((ht[f"increase_{st}"] for st in tot["included_statuses"]), D(0))
        rows = [("المجموع الكلي خارج الرسم (HT)", ht["contract"], ht["decrease"], inc_shown, tot["final_contractual_ht"])]
        if tva and "contract" in tva:
            rows.append(("الرسم على القيمة المضافة (TVA)", tva["contract"], tva["decrease"], tva["increase"], tva["final"]))
            rows.append(("المجموع الكلي بكل الرسوم (TTC)", ht["contract"] + tva["contract"], ht["decrease"] + tva["decrease"],
                         inc_shown + tva["increase"], tot["final_contractual_ttc"]))
        elif tva and "final" in tva and tot.get("final_contractual_ttc") is not None:
            # الضريبة على الصافي النهائي: لا نعرض تفصيل TVA لمكوّنات لم تُحسب منفردة.
            rows.append(("TVA على الصافي النهائي HT", None, None, None, tva["final"]))
            rows.append(("المجموع النهائي بكل الرسوم (TTC)", None, None, None, tot["final_contractual_ttc"]))
        for label, a, b, c_, f in rows:
            row = add_row(t, [""] * 12, bold=True)
            merge(row, 0, 7, label)
            for col, v in zip((8, 9, 10, 11), (a, b, c_, f)):
                _cell_text(row.cells[col], fmt(v) if v is not None else "", True, 9)
    else:
        row = add_row(t, [""] * 12, bold=True)
        merge(row, 0, 11, "الإجماليات النهائية غير محسوبة: لم تُحدَّد بعد طريقة إدراج الزيادات أو البيانات ناقصة", align=C)
    if pending_note:
        para(doc, "* كمية زيادة غير معتمدة بعد: تظهر في عمود الكمية فقط، ولا يدخل مبلغها في الإجماليات ما لم تُدرج صراحةً.", False, 8, color="B3402F")
    if ht and tva and tot.get("final_contractual_ttc") is not None:
        para(doc, "", space_after=6)
        para(doc, "الحوصلة", True, 11, C).paragraph_format.keep_with_next = True
        s = new_table(doc, 2, [10.0, 6.0])
        if "contract" in tva:
            add_row(s, ["مبلغ العقد بكل الرسوم", fmt(ht["contract"] + tva["contract"])], size=9, aligns=[R, C])
            add_row(s, ["مبلغ الأشغال غير المنجزة في إطار العقد بكل الرسوم", fmt(ht["decrease"] + tva["decrease"])], size=9, aligns=[R, C])
            add_row(s, ["مبلغ الأشغال الإضافية داخل إطار العقد بكل الرسوم", fmt(tot["final_contractual_ttc"] - (ht["contract"] + tva["contract"]) + ht["decrease"] + tva["decrease"])],
                    size=9, aligns=[R, C])
        else:
            add_row(s, ["أساس الضريبة: الصافي النهائي HT", fmt(tot["final_contractual_ht"])], size=9, aligns=[R, C])
            add_row(s, ["TVA على الصافي النهائي", fmt(tva.get("final"))], size=9, aligns=[R, C])
        add_row(s, ["مبلغ الأشغال المنجزة بكل الرسوم (TTC)", fmt(tot["final_contractual_ttc"])], bold=True, size=9, aligns=[R, C])
        keep_together(s)
    para(doc, "", space_after=8)
    _words_line(doc, "أوقف هذا الجرد على مبلغ بكامل الرسوم قدره:", tot.get("final_contractual_ttc"), state)
    ex = _extra(proj)
    _signatures(doc, ["المتعامل المتعاقد", "مكتب الدراسات", "المصلحة المتعاقدة"],
                [ex.get("manager"), None, ex.get("owner_rep") or proj.get("owner")], _place_date(d, proj))


BUILDERS = {"bayan": build_bayan, "hawsala": build_hawsala, "situation": build_situation, "inventory": build_inventory}


_PPR = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr", "suppressLineNumbers",
        "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE",
        "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents",
        "suppressOverlap", "jc", "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle",
        "rPr", "sectPr", "pPrChange"]
_TCPR = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar", "textDirection",
         "tcFitText", "vAlign", "hideMark"]


def _page_field(p, instruction):
    """حقل Word ديناميكي (PAGE / NUMPAGES) بالبنية القياسية: كل جزء في Run مستقل بنفس التنسيق."""
    def mk():
        r = p.add_run()
        rpr = r._r.get_or_add_rPr()
        _set(rpr, "w:rFonts", **{"w:ascii": FONT, "w:hAnsi": FONT, "w:cs": FONT})
        _set(rpr, "w:color", **{"w:val": "56616B"})
        _set(rpr, "w:sz", **{"w:val": "16"}); _set(rpr, "w:szCs", **{"w:val": "16"})
        _set(rpr, "w:rtl")
        _reorder(rpr, _RPR)
        return r
    for kind in ("begin", "instr", "separate", "text", "end"):
        r = mk()
        if kind == "instr":
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = f" {instruction} "
        elif kind == "text":
            el = OxmlElement("w:t"); el.text = "1"
        else:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind)
        r._r.append(el)


def _add_footer(doc, proj):
    section = doc.sections[0]
    section.footer_distance = Cm(0.65)
    p = section.footer.paragraphs[0]
    _bidi(p)
    _jc(p, C)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(0)
    # فاصل علوي خفيف للتذييل
    ppr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single"); top.set(qn("w:sz"), "4")
    top.set(qn("w:space"), "3"); top.set(qn("w:color"), "C8D0D8")
    borders.append(top); ppr.append(borders)
    project = (proj.get("name") or "وثائق الأشغال").strip()
    contract = ltr((proj.get("contract_no") or "").strip()) or ""
    run(p, project + (f" | العقد: {contract}" if contract else "") + " | صفحة ", size=8, color="56616B")
    _page_field(p, "PAGE")
    run(p, " / ", size=8, color="56616B")
    _page_field(p, "NUMPAGES")


def _rtl_defaults(doc):
    """اتجاه RTL على مستوى القسم (ترقيم الصفحات والأعمدة) وعلى النمط الافتراضي، ولغة عربية للنص المعقّد."""
    for sec in doc.sections:
        sp = sec._sectPr
        if sp.find(qn("w:bidi")) is None:
            b = OxmlElement("w:bidi")
            anchor = sp.find(qn("w:docGrid"))
            if anchor is not None:
                anchor.addprevious(b)
            else:
                sp.append(b)
    st = doc.styles.element
    dd = st.find(qn("w:docDefaults"))
    if dd is not None:
        rpr = dd.find(qn("w:rPrDefault") + "/" + qn("w:rPr"))
        if rpr is not None:
            _set(rpr, "w:rFonts", **{"w:ascii": FONT, "w:hAnsi": FONT, "w:cs": FONT})
            _set(rpr, "w:rtl")
            _set(rpr, "w:lang", **{"w:val": "fr-FR", "w:bidi": "ar-DZ"})
            _reorder(rpr, _RPR)
    normal = doc.styles["Normal"].element
    ppr = normal.get_or_add_pPr()
    if ppr.find(qn("w:bidi")) is None:
        ppr.append(OxmlElement("w:bidi"))
    _reorder(ppr, _PPR)


def _finalize(doc, project=None):
    """ترتيب عناصر OOXML (الدمج يضيف gridSpan/vMerge بعد shd) + إصلاح zoom + تذييل موحّد ومعلومات تعريف الوثيقة."""
    if project:
        _add_footer(doc, project)
    _rtl_defaults(doc)
    parts = [doc.element.body] + [sec.footer._element for sec in doc.sections]
    for root in parts:
        for ppr in root.iter(qn("w:pPr")):
            _reorder(ppr, _PPR)
        for tcpr in root.iter(qn("w:tcPr")):
            _reorder(tcpr, _TCPR)
    for z in doc.settings.element.iter(qn("w:zoom")):
        if z.get(qn("w:percent")) is None:
            z.set(qn("w:percent"), "100")
    doc.core_properties.author = (project.get("contractor") or "إدارة وثائق الأشغال") if project else "إدارة وثائق الأشغال"
    doc.core_properties.subject = "وثيقة إدارية لمتابعة الأشغال"
    doc.core_properties.language = "ar-DZ"


def export_docx(conn, doc_id, path):
    """يكتب ملفًا جديدًا في path. لا يغيّر حالة الوثيقة ولا يلمس قوالب أصلية."""
    d, state, res = load_model(conn, doc_id)
    doc = Document()
    BUILDERS[d["doc_type"]](doc, d, state, res)
    _finalize(doc, state.get("project") or {})
    doc.save(path)
    return path


def _office_executable():
    """العثور على LibreOffice في الأنظمة الشائعة من دون افتراض تثبيت في PATH."""
    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return found
    if os.name == "nt":
        candidates = [
            Path(os.environ.get("PROGRAMFILES", r"C:\Program Files")) / "LibreOffice/program/soffice.exe",
            Path(os.environ.get("PROGRAMFILES(X86)", r"C:\Program Files (x86)")) / "LibreOffice/program/soffice.exe",
        ]
        for candidate in candidates:
            if candidate.is_file():
                return str(candidate)
    return None


def export_pdf(conn, doc_id, path, timeout=90):
    """تصدير PDF عبر LibreOffice مع مجلد مؤقت مستقل. لا يعتمد على Word ولا يغير حالة الوثيقة."""
    office = _office_executable()
    if not office:
        raise RuntimeError("تصدير PDF يتطلب تثبيت LibreOffice. ثبّته ثم أعد تشغيل البرنامج، أو صدّر Word واحفظه PDF من Word.")
    target = Path(path).expanduser().resolve()
    if target.suffix.lower() != ".pdf":
        target = target.with_suffix(".pdf")
    target.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="works_pdf_") as td:
        work = Path(td)
        docx_path = work / "document.docx"
        pdf_path = work / "document.pdf"
        profile = (work / "lo_profile").resolve().as_uri()
        export_docx(conn, doc_id, str(docx_path))
        cmd = [office, "--headless", f"-env:UserInstallation={profile}", "--convert-to", "pdf", "--outdir", str(work), str(docx_path)]
        try:
            proc = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout, check=False)
        except subprocess.TimeoutExpired as exc:
            raise RuntimeError(f"انتهت مهلة تصدير PDF بعد {timeout} ثانية. أغلق LibreOffice وحاول مجددًا.") from exc
        if proc.returncode != 0 or not pdf_path.is_file() or pdf_path.stat().st_size < 500:
            details = (proc.stderr or proc.stdout or "لا توجد تفاصيل إضافية").strip()
            raise RuntimeError(f"فشل تحويل الوثيقة إلى PDF: {details}")
        shutil.copy2(pdf_path, target)
    return str(target)

