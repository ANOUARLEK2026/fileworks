import json, os, tempfile, unittest
from docx import Document
from worksapp import storage as S, services as SV, export as E
from tests.test_core import ITEMS

NB = "\u00a0"


def doc_text(path):
    d = Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts += [c.text for c in row.cells]
    # علامات LRM تُضاف حول مجموعات الآلاف لمنع انعكاسها في Word؛ لا تؤثر على النص المقروء
    return "\n".join(parts).replace("\u200e", "")


class Export(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp(); self.c = S.connect(os.path.join(self.dir, "e.db"))
        self.pid = S.create_project(self.c, "ساحة الشباب", contract_no="2026/09", contract_date="2026-02-23",
                                    contractor="الأفضل", extra={"manager": "بكاري"})
        for code, price, q in ITEMS:
            S.add_item(self.c, self.pid, code, f"بند {code}", "و", price, q, section="الأشغال")
        S.set_project_settings(self.c, self.pid, {"tva_rate": "0.19", "situation_mode": "period",
                                                  "increase_statuses": ["approved"],
                                                  "deductions": [{"code": "g", "label": "Retenue de garantie", "method": "none"}]})

    def tearDown(self):
        try: self.c.close()
        except Exception: pass

    def situation(self):
        d = SV.create_draft(self.c, self.pid, "situation")
        pl = json.loads(S.get_document(self.c, d)["payload_json"])
        for code, price, q in ITEMS:
            pl["rows"][code].update({"prev": "0", "period": str(q)})
        S.update_draft_payload(self.c, d, pl)
        return d

    def test_draft_banner_and_amounts(self):
        d = self.situation(); p = os.path.join(self.dir, "s.docx"); E.export_docx(self.c, d, p)
        t = doc_text(p)
        self.assertIn("مسودة — غير معتمدة", t)
        self.assertIn(f"9{NB}333{NB}403.84", t)
        self.assertIn("المبلغ الصافي للدفع بما فيه الرسم (بالأحرف)", t)
        self.assertEqual(S.get_document(self.c, d)["status"], "draft")        # التصدير ليس اعتمادًا

    def test_issued_export_comes_from_snapshot(self):
        d = self.situation(); SV.issue(self.c, d, explicit_approval=True)
        p1 = os.path.join(self.dir, "a.docx"); E.export_docx(self.c, d, p1)
        S.update_item_price(self.c, [i for i in S.list_items(self.c, self.pid) if i["code"] == "01"][0]["id"], 1, "x")
        S.update_project(self.c, self.pid, name="اسم مختلف تمامًا", contractor="مقاول آخر")
        p2 = os.path.join(self.dir, "b.docx"); E.export_docx(self.c, d, p2)
        self.assertEqual(doc_text(p1), doc_text(p2))
        self.assertNotIn("مسودة — غير معتمدة", doc_text(p1))
        self.assertIn("وضعية الأشغال رقم: 01", doc_text(p1))
        self.assertNotIn("اسم مختلف", doc_text(p2))

    def test_correction_label(self):
        d = self.situation(); SV.issue(self.c, d, explicit_approval=True)
        k = S.create_correction(self.c, d, "خطأ في كمية البند 03")
        p = os.path.join(self.dir, "k.docx"); E.export_docx(self.c, k, p)
        self.assertIn("نسخة تصحيحية رقم 1", doc_text(p)); self.assertIn("خطأ في كمية البند 03", doc_text(p))

    def test_all_types_export_and_reopen(self):
        d = self.situation()
        b = SV.create_draft(self.c, self.pid, "bayan")
        i1 = S.list_items(self.c, self.pid)[0]["id"]
        S.add_measure_line(self.c, i1, "length", dict(n=15, L=".25"))
        h = SV.create_draft(self.c, self.pid, "hawsala")
        pl = json.loads(S.get_document(self.c, h)["payload_json"]); pl["rows"]["01"].update({"prev": "0", "period": "15"})
        S.update_draft_payload(self.c, h, pl)
        inv = SV.create_draft(self.c, self.pid, "inventory")
        pl = json.loads(S.get_document(self.c, inv)["payload_json"])
        pl["rows"]["01"].update({"decrease": "1", "increase": "2", "increase_status": "approved"}); S.update_draft_payload(self.c, inv, pl)
        for did in (b, h, inv):
            p = os.path.join(self.dir, f"{did}.docx"); E.export_docx(self.c, did, p)
            self.assertGreater(len(Document(p).tables), 0)
        inv_text = doc_text(os.path.join(self.dir, f"{inv}.docx"))
        self.assertIn("أوقف هذا الجرد على مبلغ بكامل الرسوم قدره", inv_text)
        # عوامل القياس تُكتب 0.25 لا .25
        tb = doc_text(os.path.join(self.dir, f"{b}.docx"))
        self.assertIn("0.25", tb); self.assertNotIn(" .25", tb)

    def test_inventory_on_net_tax_basis_exports_correct_summary(self):
        ps = S.project_settings(self.c, self.pid)
        ps["tva_basis"] = "on_net"
        ps["increase_statuses"] = ["approved"]
        S.set_project_settings(self.c, self.pid, ps)
        inv = SV.create_draft(self.c, self.pid, "inventory")
        payload = json.loads(S.get_document(self.c, inv)["payload_json"])
        payload["rows"]["01"].update({"decrease": "1", "increase": "2", "increase_status": "approved"})
        S.update_draft_payload(self.c, inv, payload)
        path = os.path.join(self.dir, "inventory-on-net.docx")
        E.export_docx(self.c, inv, path)
        text = doc_text(path)
        self.assertIn("TVA على الصافي النهائي", text)
        self.assertIn("أساس الضريبة: الصافي النهائي HT", text)
        self.assertIn("أوقف هذا الجرد على مبلغ بكامل الرسوم قدره", text)
        self.assertNotIn("المجاميع النهائية غير محسوبة", text)

    def test_quantity_format_trims_unneeded_zeroes(self):
        from worksapp.qtyfmt import format_qty
        self.assertEqual(format_qty("2.000", "م²", "hawsala"), "2.00")
        self.assertEqual(format_qty("1.125", "م²", "hawsala"), "1.125")   # لا إخفاء لدقة التقريب
        self.assertEqual(format_qty("12.000", "م³", "hawsala"), "12")

    def test_footer_has_project_and_page_fields(self):
        d = self.situation(); p = os.path.join(self.dir, "footer.docx"); E.export_docx(self.c, d, p)
        doc = Document(p)
        footer_xml = doc.sections[0].footer._element.xml
        self.assertIn("ساحة الشباب", footer_xml)
        self.assertIn("PAGE", footer_xml)
        self.assertIn("NUMPAGES", footer_xml)

    @unittest.skipUnless(E._office_executable(), "LibreOffice is not installed")
    def test_pdf_export_is_valid_and_does_not_issue_document(self):
        d = self.situation(); p = os.path.join(self.dir, "document.pdf")
        saved = E.export_pdf(self.c, d, p)
        with open(saved, "rb") as fh:
            self.assertEqual(fh.read(5), b"%PDF-")
        self.assertEqual(S.get_document(self.c, d)["status"], "draft")

    def test_factor(self):
        self.assertEqual(E.factor(".6"), "0.60"); self.assertEqual(E.factor("15"), "15"); self.assertEqual(E.factor("2.155"), "2.155")


class ExportOptions(Export):
    def test_final_suffix_and_form_summary(self):
        ps = S.project_settings(self.c, self.pid); ps["summary_style"] = "form"
        ps["deductions"] = [{"code": "guarantee", "label": "Retenue de garantie", "method": "percent", "rate": "0.05",
                             "base": "ht_period", "scope": "per_situation"}]
        S.set_project_settings(self.c, self.pid, ps)
        d = self.situation()
        pl = json.loads(S.get_document(self.c, d)["payload_json"]); pl["config"] = {"final": True}; S.update_draft_payload(self.c, d, pl)
        SV.issue(self.c, d, explicit_approval=True)
        p = os.path.join(self.dir, "f.docx"); E.export_docx(self.c, d, p); t = doc_text(p)
        self.assertIn("رقم: 01 والنهائي", t)
        self.assertIn("(6) المبلغ الصافي للدفع", t)
        net = SV.preview(self.c, d)["results"]["net"] if False else json.loads(S.get_document(self.c, d)["snapshot_json"])["results"]["net"]
        self.assertIn(E.fmt(net), t)       # الصافي في الملخص = الصافي المحسوب
